#!/usr/bin/env python3
"""
Trigger Checkmarx One scan recalculations based on project priority in pr_codigo.xlsx.

Priority schedule:
  1 -> weekly   (recalculate if last scan is older than 7 days)
  2 -> biweekly (recalculate if last scan is older than 15 days)
  3 -> monthly  (recalculate if last scan is older than 30 days)
"""

from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import openpyxl
import requests
from dotenv import load_dotenv
import os

API_VERSION_HEADER = "application/json; version=1.0"
DEFAULT_PRIORITY_THRESHOLDS_DAYS = {1: 7, 2: 15, 3: 30}
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL = SCRIPT_DIR / "pr_codigo.xlsx"
DEFAULT_REPORT_DIR = SCRIPT_DIR / "reports"
VALID_ENGINES = {"sca", "sast", "kics", "containers", "system", "apisec"}
STATUS_DETAIL_SKIP = {"general"}

REPORT_COLUMNS = [
    "execution_time_utc",
    "project_name",
    "priority",
    "threshold_days",
    "status",
    "reason",
    "project_id",
    "last_scan_date",
    "days_since_last_scan",
    "branch",
    "engines",
    "new_scan_id",
    "error",
]


@dataclass
class ProjectRow:
    name: str
    priority: int


@dataclass
class CheckmarxConfig:
    tenant: str
    base_url: str
    auth_url: str
    api_key: str


@dataclass
class ExecutionResult:
    execution_time_utc: str
    project_name: str
    priority: int
    threshold_days: int
    status: str
    reason: str
    project_id: str = ""
    last_scan_date: str = ""
    days_since_last_scan: str = ""
    branch: str = ""
    engines: str = ""
    new_scan_id: str = ""
    error: str = ""

    def as_row(self) -> list[str]:
        return [str(getattr(self, column)) for column in REPORT_COLUMNS]


class CheckmarxClient:
    def __init__(self, config: CheckmarxConfig) -> None:
        self.config = config
        self.session = requests.Session()
        self._access_token: str | None = None

    def authenticate(self) -> None:
        token_url = (
            f"{self.config.auth_url.rstrip('/')}/auth/realms/"
            f"{self.config.tenant}/protocol/openid-connect/token"
        )
        response = self.session.post(
            token_url,
            data={
                "grant_type": "refresh_token",
                "client_id": "ast-app",
                "refresh_token": self.config.api_key,
            },
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            timeout=60,
        )
        if response.status_code != 200:
            raise RuntimeError(
                f"Authentication failed ({response.status_code}): {response.text}"
            )
        payload = response.json()
        self._access_token = payload["access_token"]
        self.session.headers.update(
            {
                "Authorization": f"Bearer {self._access_token}",
                "Accept": API_VERSION_HEADER,
                "Content-Type": API_VERSION_HEADER,
            }
        )

    def _api_url(self, path: str) -> str:
        return f"{self.config.base_url.rstrip('/')}{path}"

    def find_project_by_name(self, name: str) -> dict[str, Any] | None:
        response = self.session.get(
            self._api_url("/api/projects"),
            params={"name": name, "limit": 20, "offset": 0},
            timeout=60,
        )
        response.raise_for_status()
        projects = response.json().get("projects", [])
        exact_matches = [p for p in projects if p.get("name") == name]
        if not exact_matches:
            return None
        if len(exact_matches) > 1:
            raise ValueError(
                f"Multiple projects named '{name}' were found; resolve duplicates in Checkmarx One."
            )
        return exact_matches[0]

    def get_last_completed_scan(self, project_id: str) -> dict[str, Any] | None:
        response = self.session.get(
            self._api_url("/api/scans"),
            params={
                "project-id": project_id,
                "statuses": "Completed",
                "limit": 1,
                "offset": 0,
                "sort": "-created_at",
            },
            timeout=60,
        )
        response.raise_for_status()
        scans = response.json().get("scans", [])
        if not scans:
            return None
        scan = scans[0]
        if scan.get("status") != "Completed":
            return None
        return self._ensure_scan_engines(scan)

    def get_scan(self, scan_id: str) -> dict[str, Any]:
        response = self.session.get(
            self._api_url(f"/api/scans/{scan_id}"),
            timeout=60,
        )
        response.raise_for_status()
        return response.json()

    def _ensure_scan_engines(self, scan: dict[str, Any]) -> dict[str, Any]:
        if extract_engines_from_scan(scan):
            return scan
        scan_id = scan.get("id")
        if not scan_id:
            return scan
        return self.get_scan(scan_id)

    def recalculate(
        self,
        project_id: str,
        *,
        branch: str | None = None,
        engines: list[str] | None = None,
    ) -> dict[str, Any]:
        body: dict[str, Any] = {"project_id": project_id}
        if branch:
            body["branch"] = branch
        if engines:
            body["engines"] = engines
        response = self.session.post(
            self._api_url("/api/scans/recalculate"),
            json=body,
            timeout=60,
        )
        response.raise_for_status()
        if response.content:
            return response.json()
        return {}


def load_config() -> CheckmarxConfig:
    load_dotenv(SCRIPT_DIR / ".env")
    tenant = os.getenv("CHECKMARX_TENANT", "").strip()
    base_url = os.getenv("CHECKMARX_BASE_URL", "").strip()
    auth_url = os.getenv("CHECKMARX_AUTH_URL", "").strip()
    api_key = os.getenv("CHECKMARX_API_KEY", "").strip()

    missing = [
        name
        for name, value in [
            ("CHECKMARX_TENANT", tenant),
            ("CHECKMARX_BASE_URL", base_url),
            ("CHECKMARX_AUTH_URL", auth_url),
            ("CHECKMARX_API_KEY", api_key),
        ]
        if not value
    ]
    if missing:
        raise RuntimeError(
            "Missing required environment variables in .env: "
            + ", ".join(missing)
        )

    return CheckmarxConfig(
        tenant=tenant,
        base_url=base_url,
        auth_url=auth_url,
        api_key=api_key,
    )


def load_priority_thresholds() -> dict[int, int]:
    thresholds: dict[int, int] = {}
    for priority, default_days in DEFAULT_PRIORITY_THRESHOLDS_DAYS.items():
        raw = os.getenv(f"PRIORITY_{priority}_DAYS", str(default_days)).strip()
        try:
            days = int(raw)
        except ValueError as exc:
            raise RuntimeError(
                f"PRIORITY_{priority}_DAYS must be a whole number, got: {raw!r}"
            ) from exc
        if days < 1:
            raise RuntimeError(f"PRIORITY_{priority}_DAYS must be at least 1, got: {days}")
        thresholds[priority] = days
    return thresholds


def extract_engines_from_scan(scan: dict[str, Any]) -> list[str]:
    engines = scan.get("engines")
    if isinstance(engines, list) and engines:
        return [e for e in engines if isinstance(e, str) and e in VALID_ENGINES]

    metadata = scan.get("metadata") or {}
    configs = metadata.get("configs") or []
    from_configs = [
        cfg["type"]
        for cfg in configs
        if isinstance(cfg, dict) and cfg.get("type") in VALID_ENGINES
    ]
    if from_configs:
        return from_configs

    from_status = [
        detail["name"]
        for detail in scan.get("statusDetails") or []
        if isinstance(detail, dict)
        and detail.get("name") not in STATUS_DETAIL_SKIP
        and detail.get("name") in VALID_ENGINES
    ]
    return from_status


def parse_scan_timestamp(scan: dict[str, Any]) -> datetime:
    raw = scan.get("updatedAt") or scan.get("createdAt")
    if not raw:
        raise ValueError("Scan is missing createdAt/updatedAt.")
    normalized = raw.replace("Z", "+00:00")
    parsed = datetime.fromisoformat(normalized)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def load_projects_from_excel(
    path: Path,
    valid_priorities: set[int],
) -> list[ProjectRow]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    try:
        name_idx = header.index("project name")
        priority_idx = header.index("priority")
    except ValueError as exc:
        raise ValueError(
            "Excel file must contain 'Project name' and 'Priority' columns."
        ) from exc

    projects: list[ProjectRow] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        name = str(row[name_idx]).strip()
        if not name:
            print(f"[WARN] Row {row_number}: empty project name, skipping.")
            continue
        try:
            priority = int(row[priority_idx])
        except (TypeError, ValueError):
            print(f"[WARN] Row {row_number}: invalid priority for '{name}', skipping.")
            continue
        if priority not in valid_priorities:
            print(
                f"[WARN] Row {row_number}: unsupported priority {priority} for '{name}', skipping."
            )
            continue
        projects.append(ProjectRow(name=name, priority=priority))
    return projects


def should_recalculate(
    last_scan: dict[str, Any] | None,
    threshold_days: int,
    now: datetime,
) -> tuple[bool, str]:
    if last_scan is None:
        return False, "no completed scan found"

    last_scan_at = parse_scan_timestamp(last_scan)
    age = now - last_scan_at
    if age > timedelta(days=threshold_days):
        return True, (
            f"last scan on {last_scan_at.date().isoformat()} "
            f"({age.days} days ago, threshold {threshold_days} days)"
        )
    return False, (
        f"last scan on {last_scan_at.date().isoformat()} "
        f"({age.days} days ago, threshold {threshold_days} days)"
    )


def scan_age_fields(last_scan: dict[str, Any] | None, now: datetime) -> tuple[str, str]:
    if last_scan is None:
        return "", ""
    last_scan_at = parse_scan_timestamp(last_scan)
    return last_scan_at.date().isoformat(), str((now - last_scan_at).days)


def write_csv_report(path: Path, results: list[ExecutionResult]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(REPORT_COLUMNS)
        for result in results:
            writer.writerow(result.as_row())


def write_excel_report(
    path: Path,
    results: list[ExecutionResult],
    *,
    summary: dict[str, str],
) -> None:
    workbook = openpyxl.Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Field", "Value"])
    for key, value in summary.items():
        summary_sheet.append([key, value])

    results_sheet = workbook.create_sheet("Results")
    results_sheet.append(REPORT_COLUMNS)
    for result in results:
        results_sheet.append(result.as_row())

    workbook.save(path)


def write_reports(
    report_dir: Path,
    results: list[ExecutionResult],
    *,
    summary: dict[str, str],
    run_timestamp: datetime,
) -> tuple[Path, Path]:
    report_dir.mkdir(parents=True, exist_ok=True)
    stamp = run_timestamp.strftime("%Y-%m-%d_%H%M%S")
    csv_path = report_dir / f"rescans_report_{stamp}.csv"
    xlsx_path = report_dir / f"rescans_report_{stamp}.xlsx"
    write_csv_report(csv_path, results)
    write_excel_report(xlsx_path, results, summary=summary)
    return csv_path, xlsx_path


def process_project(
    client: CheckmarxClient,
    project: ProjectRow,
    *,
    now: datetime,
    dry_run: bool,
    priority_thresholds: dict[int, int],
) -> ExecutionResult:
    execution_time = now.isoformat()
    threshold_days = priority_thresholds[project.priority]
    base = ExecutionResult(
        execution_time_utc=execution_time,
        project_name=project.name,
        priority=project.priority,
        threshold_days=threshold_days,
        status="",
        reason="",
    )

    try:
        cx_project = client.find_project_by_name(project.name)
        if cx_project is None:
            base.status = "SKIPPED"
            base.reason = "project not found in Checkmarx One"
            return base

        project_id = cx_project["id"]
        base.project_id = project_id
        last_scan = client.get_last_completed_scan(project_id)
        last_scan_date, days_since = scan_age_fields(last_scan, now)
        base.last_scan_date = last_scan_date
        base.days_since_last_scan = days_since
        if last_scan:
            base.branch = last_scan.get("branch") or ""

        do_recalc, reason = should_recalculate(last_scan, threshold_days, now)
        if not do_recalc:
            base.status = "SKIPPED"
            base.reason = reason
            return base

        engines = extract_engines_from_scan(last_scan)
        base.engines = ", ".join(engines)
        if not engines:
            base.status = "SKIPPED"
            base.reason = "last completed scan has no engines to recalculate"
            return base

        base.reason = reason
        if dry_run:
            base.status = "DRY RUN"
            return base

        result = client.recalculate(
            project_id,
            branch=last_scan.get("branch"),
            engines=engines,
        )
        base.status = "TRIGGERED"
        base.new_scan_id = result.get("id", "")
        return base
    except requests.HTTPError as exc:
        body = exc.response.text if exc.response is not None else str(exc)
        status_code = exc.response.status_code if exc.response else "?"
        base.status = "ERROR"
        base.reason = "API request failed"
        base.error = f"HTTP {status_code}: {body}"
        return base
    except Exception as exc:
        base.status = "ERROR"
        base.reason = "unexpected error"
        base.error = str(exc)
        return base


def run(excel_path: Path, dry_run: bool, report_dir: Path) -> int:
    config = load_config()
    priority_thresholds = load_priority_thresholds()
    projects = load_projects_from_excel(excel_path, set(priority_thresholds))
    if not projects:
        print("No projects found in the Excel file.")
        return 1

    client = CheckmarxClient(config)
    print("Authenticating with Checkmarx One...")
    client.authenticate()
    print("Authentication successful.")
    print(
        "Priority thresholds (days): "
        + ", ".join(f"{p}={d}" for p, d in sorted(priority_thresholds.items()))
        + "\n"
    )

    now = datetime.now(timezone.utc)
    results: list[ExecutionResult] = []
    triggered = 0
    skipped = 0
    errors = 0

    for project in projects:
        print(f"Project: {project.name} (priority {project.priority})")
        result = process_project(
            client,
            project,
            now=now,
            dry_run=dry_run,
            priority_thresholds=priority_thresholds,
        )
        results.append(result)

        if result.status == "TRIGGERED":
            print(f"  [RECALC] {result.reason} (engines: {result.engines}).")
            print(f"  [OK] Recalculation started (scan id: {result.new_scan_id}).\n")
            triggered += 1
        elif result.status == "DRY RUN":
            print(f"  [RECALC] {result.reason} (engines: {result.engines}).")
            print("  [DRY RUN] Would call POST /api/scans/recalculate.\n")
            triggered += 1
        elif result.status == "ERROR":
            print(f"  [ERROR] {result.error}\n")
            errors += 1
        else:
            print(f"  [SKIP] {result.reason}.\n")
            skipped += 1

    summary = {
        "execution_time_utc": now.isoformat(),
        "tenant": config.tenant,
        "checkmarx_base_url": config.base_url,
        "input_file": str(excel_path),
        "dry_run": str(dry_run),
        "priority_1_days": str(priority_thresholds[1]),
        "priority_2_days": str(priority_thresholds[2]),
        "priority_3_days": str(priority_thresholds[3]),
        "total_projects": str(len(results)),
        "triggered": str(triggered),
        "skipped": str(skipped),
        "errors": str(errors),
    }
    csv_path, xlsx_path = write_reports(
        report_dir,
        results,
        summary=summary,
        run_timestamp=now,
    )

    print(
        f"Done. triggered={triggered}, skipped={skipped}, errors={errors}"
        + (" (dry run)" if dry_run else "")
    )
    print(f"Report written to:\n  {csv_path}\n  {xlsx_path}")
    return 1 if errors else 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run Checkmarx One scan recalculations based on project priority."
    )
    parser.add_argument(
        "--excel",
        type=Path,
        default=Path(os.getenv("PROJECTS_EXCEL", DEFAULT_EXCEL)),
        help="Path to the Excel file with Project name and Priority columns.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Report actions without calling the recalculate API.",
    )
    parser.add_argument(
        "--report-dir",
        type=Path,
        default=Path(os.getenv("REPORT_DIR", DEFAULT_REPORT_DIR)),
        help="Directory where CSV and Excel execution reports are saved.",
    )
    args = parser.parse_args()

    if not args.excel.exists():
        print(f"Excel file not found: {args.excel}", file=sys.stderr)
        return 1

    try:
        return run(args.excel, dry_run=args.dry_run, report_dir=args.report_dir)
    except Exception as exc:
        print(f"Fatal error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
